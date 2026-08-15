from docx import Document


class DOCXParser:

    def extract_text(self, file):
        document = Document(file)

        text = []

     
        for paragraph in document.paragraphs:
            paragraph_text = paragraph.text.strip()

            if paragraph_text:
                text.append(paragraph_text)

      
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    text.append(" | ".join(row_text))

        return "\n".join(text).strip()